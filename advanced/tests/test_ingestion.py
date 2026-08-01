"""
Minimal test coverage for the ingestion layer — required for the CI/CD deliverable.
Tests use small, generated sample files rather than committed binary fixtures,
so the test suite doesn't reintroduce the large-file problem we already hit once.
"""

import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parents[2]))

import pytest
from docx import Document as DocxDocument
import openpyxl

from advanced.ingestion.loaders import (
    load_document, load_txt, load_docx, load_xlsx,
    UnsupportedFileTypeError, DocumentLoadError
)
from advanced.ingestion.chunking import chunk_documents


@pytest.fixture
def tmp_txt_file(tmp_path):
    f = tmp_path / "sample.txt"
    f.write_text("This is a test document. It has multiple sentences. "
                 "Used to verify the TXT loader works correctly.", encoding="utf-8")
    return f


@pytest.fixture
def tmp_empty_txt_file(tmp_path):
    f = tmp_path / "empty.txt"
    f.write_text("", encoding="utf-8")
    return f


@pytest.fixture
def tmp_docx_file(tmp_path):
    f = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("This is a test DOCX document for ingestion testing.")
    doc.save(f)
    return f


@pytest.fixture
def tmp_xlsx_file(tmp_path):
    f = tmp_path / "sample.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Score"])
    ws.append(["Alice", 90])
    ws.append(["Bob", 85])
    wb.save(f)
    return f


def test_load_txt_success(tmp_txt_file):
    docs = load_document(tmp_txt_file)
    assert len(docs) > 0
    assert "test document" in docs[0].page_content
    assert docs[0].metadata["source_file"] == "sample.txt"


def test_load_empty_txt_raises(tmp_empty_txt_file):
    with pytest.raises(DocumentLoadError):
        load_document(tmp_empty_txt_file)


def test_load_docx_success(tmp_docx_file):
    docs = load_document(tmp_docx_file)
    assert len(docs) > 0
    assert "test DOCX document" in docs[0].page_content


def test_load_xlsx_success(tmp_xlsx_file):
    docs = load_document(tmp_xlsx_file)
    assert len(docs) > 0
    assert "Alice" in docs[0].page_content
    assert "Score" in docs[0].page_content


def test_unsupported_extension_raises(tmp_path):
    f = tmp_path / "sample.xyz"
    f.write_text("irrelevant content")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(f)


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_document(tmp_path / "does_not_exist.txt")


def test_chunking_splits_long_document(tmp_path):
    f = tmp_path / "long.txt"
    f.write_text(" ".join(["word"] * 2000), encoding="utf-8")  # long enough to force multiple chunks
    docs = load_document(f)
    chunks = chunk_documents(docs)
    assert len(chunks) > 1
    assert all(c.metadata.get("chunk_index") for c in chunks)


def test_chunking_preserves_source_metadata(tmp_txt_file):
    docs = load_document(tmp_txt_file)
    chunks = chunk_documents(docs)
    assert chunks[0].metadata["source_file"] == "sample.txt"